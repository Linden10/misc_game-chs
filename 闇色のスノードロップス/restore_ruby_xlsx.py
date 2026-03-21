from __future__ import annotations

import argparse
import shutil
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("openpyxl is required. Install it with: pip install openpyxl") from exc

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


def soft_index(line_no: int, start: int, end: int) -> int:
    return line_no * 1000000 + start * 1000 + end


def decode_soft_index(index: int) -> tuple[int, int, int]:
    return index // 1000000, (index // 1000) % 1000, index % 1000


def normalize_text(text: object) -> str:
    if text is None:
        return ""
    normalized = unicodedata.normalize("NFKC", str(text).replace("\r", "").replace("\n", ""))
    chars: list[str] = []
    for char in normalized:
        code = ord(char)
        if 0x30A1 <= code <= 0x30F6:
            chars.append(chr(code - 0x60))
        else:
            chars.append(char)
    return "".join("".join(chars).split())


def smart_wrap(text: str, limit: int = 53) -> str:
    if not text:
        return ""
    paragraphs = str(text).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    wrapped_parts: list[str] = []
    for paragraph in paragraphs:
        if len(paragraph) <= limit:
            wrapped_parts.append(paragraph)
            continue
        remaining = paragraph
        lines: list[str] = []
        while len(remaining) > limit:
            cut = limit
            if cut < len(remaining):
                left_space = remaining.rfind(" ", 0, cut + 1)
                if left_space > 0:
                    cut = left_space + 1
                elif remaining[cut] == " ":
                    cut += 1
            lines.append(remaining[:cut].rstrip())
            remaining = remaining[cut:].lstrip()
        lines.append(remaining)
        wrapped_parts.append("\n".join(lines))
    return "\n".join(wrapped_parts)


def choose_split_point(text: str, start: int, desired: int, min_tail: int) -> int:
    max_cut = max(start + 1, len(text) - min_tail)
    desired = max(start + 1, min(desired, max_cut))
    left_space = text.rfind(" ", start + 1, desired + 1)
    if left_space != -1:
        return left_space + 1
    if desired < len(text) and text[desired] == " ":
        return desired + 1
    right_space = text.find(" ", desired, max_cut + 1)
    if right_space != -1:
        return right_space + 1
    return desired


def split_translation(text: str, count: int, weights: list[int]) -> list[str]:
    text = "" if text is None else str(text)
    if count <= 0:
        return []
    if count == 1:
        return [smart_wrap(text)]
    explicit_parts = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if len(explicit_parts) == count:
        return [smart_wrap(part) for part in explicit_parts]
    total_weight = sum(weights) or count
    cuts: list[int] = []
    start = 0
    for i in range(count - 1):
        target = round(len(text) * sum(weights[: i + 1]) / total_weight)
        cut = choose_split_point(text, start, target, min_tail=count - i - 1)
        if cut <= start:
            cut = min(len(text), start + max(1, len(text[start:]) // (count - i)))
        cuts.append(cut)
        start = cut
    parts: list[str] = []
    prev = 0
    for cut in cuts + [len(text)]:
        parts.append(smart_wrap(text[prev:cut]))
        prev = cut
    while len(parts) < count:
        parts.append("")
    return parts


def split_delimited_translation(text: str, count: int) -> list[str] | None:
    text = "" if text is None else str(text)
    if count <= 0:
        return []
    parts = [part.strip() for part in text.split("|")]
    if len(parts) != count:
        return None
    return [smart_wrap(part) for part in parts]


@dataclass
class ExtractedToken:
    index: int
    text: str
    start: int
    end: int
    line_no: int
    kind: str = "plain"


@dataclass
class SheetRow:
    excel_row: int
    index: int | None
    jp: str
    translation: str
    kind: str
    line_no: int | None


def extract_tokens_from_line(line: str, line_no: int, in_block_comment: bool) -> tuple[list[ExtractedToken], bool]:
    tokens: list[ExtractedToken] = []
    is_text = 0
    is_at = False
    is_name = False
    start = 0
    for n, char in enumerate(line):
        if char == "/":
            if in_block_comment:
                if n > 0 and line[n - 1] == "*":
                    in_block_comment = False
            else:
                if line[n + 1 : n + 2] == "*":
                    in_block_comment = True
                elif line[n + 1 : n + 2] == "/" and is_text != 2:
                    break
        if in_block_comment or (is_text == 2 and char != '"'):
            continue
        if char == "#" or char == ";":
            break
        if char == "\t" or char == "\n":
            is_at = False
            continue
        if char == "@" and line[n + 1 : n + 5] != "ruby":
            is_at = True
        if is_at and char in {" ", ",", "(", ")", "【", "「"}:
            is_at = False
        if char == '"':
            if is_text == 2:
                is_text = 0
                if line[start:n] != "":
                    tokens.append(ExtractedToken(soft_index(line_no, start, n - 1), line[start:n], start, n - 1, line_no))
            else:
                if is_text == 1:
                    tokens.append(ExtractedToken(soft_index(line_no, start, n - 1), line[start:n], start, n - 1, line_no))
                is_text = 2
                start = n + 1
        elif unicodedata.east_asian_width(char) != "Na":
            if not is_text and not is_at:
                if char == "【":
                    start = n + 1
                    is_name = True
                    is_text = 1
                else:
                    start = n
                    is_name = False
                    is_text = 1
            elif is_text == 1 and not is_at and is_name and char == "】":
                is_name = False
                is_text = 0
                tokens.append(ExtractedToken(soft_index(line_no, start, n - 1), line[start:n], start, n - 1, line_no))
        elif is_text == 1:
            is_text = 0
            tokens.append(ExtractedToken(soft_index(line_no, start, n - 1), line[start:n], start, n - 1, line_no))
    if is_text:
        end = len(line.rstrip("\n")) - 1
        if end >= start:
            tokens.append(ExtractedToken(soft_index(line_no, start, end), line[start : end + 1], start, end, line_no))
    return tokens, in_block_comment


def ruby_spans_from_line(line: str) -> list[tuple[tuple[int, int], tuple[int, int]]]:
    spans: list[tuple[tuple[int, int], tuple[int, int]]] = []
    n = 0
    in_string = False
    while n < len(line):
        if not in_string and line.startswith("//", n):
            break
        if not in_string and line.startswith("/*", n):
            end = line.find("*/", n + 2)
            if end == -1:
                break
            n = end + 2
            continue
        if line[n] == '"':
            in_string = not in_string
            n += 1
            continue
        if in_string:
            n += 1
            continue
        if line.startswith("ruby(", n):
            n += 5
            while n < len(line) and line[n].isspace():
                n += 1
            if n >= len(line):
                break
            if line[n] == '"':
                read_start = n + 1
                n += 1
                read_end = line.find('"', n)
                if read_end == -1:
                    break
                reading = (read_start, read_end - 1)
                n = read_end + 1
            else:
                read_start = n
                read_end = line.find(")", n)
                if read_end == -1:
                    break
                reading = (read_start, read_end - 1)
                n = read_end
            if n < len(line) and line[n] == ")":
                n += 1
            while n < len(line) and line[n].isspace():
                n += 1
            text_start = n
            close = line.find("ruby", n)
            if close == -1:
                break
            text_end = close - 1
            spans.append((reading, (text_start, text_end)))
            n = close + 4
            continue
        n += 1
    return spans


def parse_source_tokens(ss_path: Path) -> dict[int, ExtractedToken]:
    tokens_by_index: dict[int, ExtractedToken] = {}
    in_block_comment = False
    token_order = 0
    with ss_path.open("r", encoding="cp932", errors="ignore") as handle:
        for line_no, line in enumerate(handle):
            tokens, in_block_comment = extract_tokens_from_line(line, line_no, in_block_comment)
            spans = ruby_spans_from_line(line)
            for token in tokens:
                for reading_span, text_span in spans:
                    if reading_span[0] <= token.start and token.end <= reading_span[1]:
                        token.kind = "ruby_reading"
                        break
                    if text_span[0] <= token.start and token.end <= text_span[1]:
                        token.kind = "ruby_text"
                        break
                token.index = token_order
                tokens_by_index[token_order] = token
                token_order += 1
    return tokens_by_index


def load_original_rows(ws: Worksheet, token_map: dict[int, ExtractedToken]) -> list[SheetRow]:
    rows: list[SheetRow] = []
    for row_idx in range(2, ws.max_row + 1):
        index_value = ws.cell(row_idx, 3).value
        if index_value is None:
            continue
        try:
            index = int(index_value)
        except ValueError:
            continue
        jp = "" if ws.cell(row_idx, 1).value is None else str(ws.cell(row_idx, 1).value)
        translation = "" if ws.cell(row_idx, 2).value is None else str(ws.cell(row_idx, 2).value)
        rows.append(SheetRow(row_idx, index, jp, translation, "plain", None))
    return rows


def load_translated_rows(ws: Worksheet) -> list[SheetRow]:
    rows: list[SheetRow] = []
    for row_idx in range(2, ws.max_row + 1):
        key = ws.cell(row_idx, 1).value
        value = ws.cell(row_idx, 2).value
        if key is None and value is None:
            continue
        rows.append(
            SheetRow(
                excel_row=row_idx,
                index=None,
                jp="" if key is None else str(key),
                translation="" if value is None else str(value),
                kind="plain",
                line_no=None,
            )
        )
    return rows


def align_original_rows(rows: list[SheetRow], token_map: dict[int, ExtractedToken]) -> None:
    tokens = [token_map[idx] for idx in sorted(token_map)]
    token_pos = 0
    for row in rows:
        row_text = normalize_text(row.jp)
        if not row_text:
            continue
        while token_pos < len(tokens) and normalize_text(tokens[token_pos].text) != row_text:
            found = None
            for idx in range(token_pos + 1, min(len(tokens), token_pos + 40)):
                if normalize_text(tokens[idx].text) == row_text:
                    found = idx
                    break
            if found is None:
                break
            token_pos = found
        if token_pos >= len(tokens) or normalize_text(tokens[token_pos].text) != row_text:
            continue
        token = tokens[token_pos]
        row.kind = token.kind
        row.line_no = token.line_no
        token_pos += 1


def build_groups(rows: list[SheetRow]) -> dict[int, dict[str, list[SheetRow]]]:
    groups: dict[int, dict[str, list[SheetRow]]] = {}
    for row in rows:
        if row.line_no is None:
            continue
        group = groups.setdefault(row.line_no, {"all": [], "plain": [], "ruby_text": [], "ruby_reading": []})
        group["all"].append(row)
        group[row.kind].append(row)
    return groups


def translated_sheet_map(workbook: Workbook) -> dict[str, Worksheet]:
    return {sheet.title: sheet for sheet in workbook.worksheets}


def find_matching_sheet(original_sheet: Worksheet, translated_workbook: Workbook) -> Worksheet | None:
    by_title = translated_sheet_map(translated_workbook)
    if original_sheet.title in by_title:
        return by_title[original_sheet.title]
    if len(translated_workbook.worksheets) == 1:
        return translated_workbook.worksheets[0]
    original_name = original_sheet.title
    for sheet in translated_workbook.worksheets:
        if sheet.title == original_name:
            return sheet
    original_index = original_sheet.parent.worksheets.index(original_sheet)
    if original_index < len(translated_workbook.worksheets):
        return translated_workbook.worksheets[original_index]
    return None


def process_sheet(original_ws: Worksheet, translated_ws: Worksheet | None, token_map: dict[int, ExtractedToken]) -> None:
    original_rows = load_original_rows(original_ws, token_map)
    align_original_rows(original_rows, token_map)
    translated_rows = load_translated_rows(translated_ws) if translated_ws is not None else []
    groups = build_groups(original_rows)
    output: dict[int, str] = {row.excel_row: row.translation for row in original_rows}

    oi = 0
    ti = 0
    while oi < len(original_rows):
        row = original_rows[oi]
        next_oi = oi + 1
        group = groups.get(row.line_no)
        if group and group["ruby_text"] and group["all"] and row.index == group["all"][0].index:
            group_rows = list(group["all"])
            if ti < len(translated_rows) and normalize_text(translated_rows[ti].jp) == normalize_text(group_rows[0].jp):
                output[group_rows[0].excel_row] = smart_wrap(translated_rows[ti].translation)
                ti += 1
                group_rows = group_rows[1:]
            visible_rows = [item for item in group_rows if item.kind != "ruby_reading"]
            reading_rows = [item for item in group_rows if item.kind == "ruby_reading"]
            visible_jp = "".join(item.jp for item in visible_rows)
            if ti < len(translated_rows) and normalize_text(translated_rows[ti].jp) == normalize_text(visible_jp):
                parts = split_translation(
                    translated_rows[ti].translation,
                    len(visible_rows),
                    [max(len(item.jp), 1) for item in visible_rows],
                )
                for item, part in zip(visible_rows, parts):
                    output[item.excel_row] = part
                ti += 1
                reading_key = " | ".join(item.jp for item in reading_rows)
                if reading_rows and ti < len(translated_rows) and normalize_text(translated_rows[ti].jp) == normalize_text(reading_key):
                    reading_parts = split_delimited_translation(translated_rows[ti].translation, len(reading_rows))
                    if reading_parts is None:
                        reading_parts = split_translation(
                            translated_rows[ti].translation,
                            len(reading_rows),
                            [max(len(item.jp), 1) for item in reading_rows],
                        )
                    for item, part in zip(reading_rows, reading_parts):
                        output[item.excel_row] = part
                    ti += 1
                else:
                    for item in reading_rows:
                        if ti < len(translated_rows) and normalize_text(translated_rows[ti].jp) == normalize_text(item.jp):
                            output[item.excel_row] = smart_wrap(translated_rows[ti].translation)
                            ti += 1
                next_oi = oi + len(group["all"])
        elif ti < len(translated_rows):
            translated = translated_rows[ti]
            if normalize_text(translated.jp) == normalize_text(row.jp):
                output[row.excel_row] = smart_wrap(translated.translation)
                ti += 1
        oi = next_oi

    for row in original_rows:
        original_ws.cell(row.excel_row, 2).value = output[row.excel_row]


def workbook_source_name(workbook_path: Path, workbook: Workbook) -> str:
    return workbook_path.stem


def resolve_input_path(path: Path, temp_root: Path) -> Path:
    if path.is_dir():
        return path
    if path.suffix.lower() != ".zip":
        raise FileNotFoundError(f"Expected a directory or zip file: {path}")
    extract_dir = temp_root / path.stem
    if extract_dir.exists():
        shutil.rmtree(extract_dir)
    extract_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as archive:
        archive.extractall(extract_dir)
    return extract_dir


def collect_xlsx_files(path: Path) -> dict[str, Path]:
    return {file.name: file for file in sorted(path.rglob("*.xlsx"))}


def collect_ss_files(path: Path) -> dict[str, Path]:
    files: dict[str, Path] = {}
    for file in sorted(path.rglob("*.ss")):
        files[file.name] = file
        files[file.stem] = file
    return files


def zip_directory(folder: Path, out_zip: Path) -> None:
    if out_zip.exists():
        out_zip.unlink()
    with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as archive:
        for file in sorted(folder.rglob("*")):
            if file.is_file():
                archive.write(file, file.relative_to(folder))


def process_workbooks(original_dir: Path, translated_dir: Path, ss_dir: Path, output_dir: Path, output_zip: Path | None) -> None:
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    original_files = collect_xlsx_files(original_dir)
    translated_files = collect_xlsx_files(translated_dir)
    ss_files = collect_ss_files(ss_dir)

    for file_name, original_path in original_files.items():
        if file_name not in translated_files:
            continue
        translated_path = translated_files[file_name]
        original_book = openpyxl.load_workbook(original_path)
        translated_book = openpyxl.load_workbook(translated_path)
        source_name = workbook_source_name(original_path, original_book)
        ss_path = ss_files.get(source_name) or ss_files.get(f"{source_name}.ss")
        token_map = parse_source_tokens(ss_path) if ss_path else {}
        for original_ws in original_book.worksheets:
            translated_ws = find_matching_sheet(original_ws, translated_book)
            process_sheet(original_ws, translated_ws, token_map)
        out_path = output_dir / file_name
        original_book.save(out_path)

    if output_zip is not None:
        zip_directory(output_dir, output_zip)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--original", required=True, help="Original JP xlsx directory or zip")
    parser.add_argument("--translated", required=True, help="Translated combined xlsx directory or zip")
    parser.add_argument("--ss", required=True, help="Raw ss directory or zip")
    parser.add_argument("--output-dir", required=True, help="Directory for processed xlsx files")
    parser.add_argument("--output-zip", help="Optional zip file to create from the processed output directory")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    with tempfile.TemporaryDirectory(prefix="restore-ruby-xlsx-") as temp_dir:
        temp_root = Path(temp_dir)
        original_dir = resolve_input_path(Path(args.original).resolve(), temp_root)
        translated_dir = resolve_input_path(Path(args.translated).resolve(), temp_root)
        ss_dir = resolve_input_path(Path(args.ss).resolve(), temp_root)
        process_workbooks(
            original_dir=original_dir,
            translated_dir=translated_dir,
            ss_dir=ss_dir,
            output_dir=Path(args.output_dir).resolve(),
            output_zip=Path(args.output_zip).resolve() if args.output_zip else None,
        )


if __name__ == "__main__":
    main()
