#!/usr/bin/env python3

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


RUBY_JOINER = " | "
MAP_DIRNAME = ".ruby_flatten_maps"
DEFAULT_GAME_DIR = Path("/workspaces/misc_game-chs/Chicchakunai Mon!")
SCRIPT_DIRNAME = "ss_20250511_111752"
TEXT_DIRNAME = "Text"
TRANSLATED_TEXT_DIRNAME = "Translated_Text"
SPEAKER_NAME_MAP = {
    "ノノ": "のの",
}
VISIBLE_COMMAND_PREFIX_RE = re.compile(r'^(?:[A-Z_]+\([^)]*\))*')
SCRIPT_CONTROL_RE = re.compile(
    r'^(?:'
    r'if\b|else\b|switch\b|case\b|default\b|break\b|command\b|farcall\b|set_title\b|returnmenu\b'
    r')'
)
ASCII_LABEL_RE = re.compile(r'^[A-Za-z_][A-Za-z0-9_]*$')
RUBY_RE = re.compile(r'ruby\("([^"]*)"\)(.*?)ruby')
TOKEN_RE = re.compile(r"[A-Za-z0-9']+")

# Some translated helper rows do not repeat the exact visible English term used in
# the combined line. These overrides tell restore which visible English segment to
# carve out of the combined translation for the base-text row.
RESTORE_BASE_TEXT_OVERRIDES = {
    ("sc011_１日目前半.ss.xlsx", 10): ["Youjou"],
    ("sc011_１日目前半.ss.xlsx", 406): ["contagious"],
    ("sc012_１日目後半.ss.xlsx", 463): ["bii"],
    ("sc012_１日目後半.ss.xlsx", 493): ["C"],
    ("sc012_１日目後半.ss.xlsx", 1154): ["flitting about"],
    ("sc012_１日目後半.ss.xlsx", 1291): ["elegant"],
    ("sc021_２日目.ss.xlsx", 3427): ["accumulated"],
    ("sc031_３日目.ss.xlsx", 281): ["maybe"],
    ("sc112_ノノ２：恋人化.ss.xlsx", 19): ["contagious"],
    ("sc112_ノノ２：恋人化.ss.xlsx", 785): ["trials and hardships"],
}


@dataclass
class Segment:
    type: str
    text: str
    emitted: bool


@dataclass
class RubyGroup:
    original_start: int
    original_length: int
    flat_start: int
    combined_text: str
    ruby_text: str
    ruby_count: int
    original_keys: list[str]
    segments: list[Segment]


@dataclass
class SceneMetadata:
    version: int
    workbook: str
    sheet_name: str
    original_row_count: int
    flattened_row_count: int
    ruby_joiner: str
    groups: list[RubyGroup]


@dataclass
class SheetLayout:
    key_col: int
    translation_col: int
    index_col: int | None
    width: int


class ValidationError(RuntimeError):
    pass


def normalize_speaker_name(speaker: str) -> str:
    if speaker.endswith("？"):
        return "？？？"
    return SPEAKER_NAME_MAP.get(speaker, speaker)


def is_visible_line(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    if stripped[0] in "{}$":
        return False
    if SCRIPT_CONTROL_RE.match(stripped):
        return False
    if ASCII_LABEL_RE.match(stripped):
        return False
    return not (
        stripped.startswith("//")
        or stripped.startswith(";")
        or stripped.startswith("@")
        or stripped.startswith("#")
    )


def strip_line_control(line: str) -> str:
    if "//" in line:
        line = line.split("//", 1)[0]
    line = line.rstrip()
    if line.endswith("R"):
        line = line[:-1]
    return VISIBLE_COMMAND_PREFIX_RE.sub("", line, count=1)


def parse_body_segments(body: str) -> list[Segment]:
    segments: list[Segment] = []
    cursor = 0
    for match in RUBY_RE.finditer(body):
        plain = body[cursor:match.start()]
        segments.append(Segment("plain", plain, bool(plain.strip())))
        segments.append(Segment("ruby", match.group(1), True))
        segments.append(Segment("base", match.group(2), True))
        cursor = match.end()
    segments.append(Segment("plain", body[cursor:], bool(body[cursor:].strip())))
    return segments


def parse_scene(scene_path: Path) -> tuple[list[str], list[RubyGroup]]:
    expected_rows: list[str] = []
    groups: list[RubyGroup] = []

    for raw_line in scene_path.read_text(encoding="cp932").splitlines():
        if not is_visible_line(raw_line):
            continue

        line = strip_line_control(raw_line)
        if line.startswith("【") and "】" in line:
            speaker, line = line[1:].split("】", 1)
            expected_rows.append(normalize_speaker_name(speaker))

        if "ruby(" not in line:
            if line.strip():
                expected_rows.append(line)
            continue

        segments = parse_body_segments(line)
        original_keys = [segment.text for segment in segments if segment.emitted]
        if not original_keys:
            continue

        group = RubyGroup(
            original_start=len(expected_rows) + 1,
            original_length=len(original_keys),
            flat_start=0,
            combined_text="".join(
                segment.text
                for segment in segments
                if segment.type != "ruby" and segment.emitted
            ),
            ruby_text=RUBY_JOINER.join(
                segment.text for segment in segments if segment.type == "ruby"
            ),
            ruby_count=sum(1 for segment in segments if segment.type == "ruby"),
            original_keys=original_keys,
            segments=segments,
        )
        groups.append(group)
        expected_rows.extend(original_keys)

    return expected_rows, groups


def locate_ruby_groups(
    sheet_rows: list[list[Any]],
    groups: list[RubyGroup],
    scene_name: str,
) -> list[RubyGroup]:
    actual_keys = [row[0] for row in sheet_rows]
    located_groups: list[RubyGroup] = []
    search_start = 1

    for group in groups:
        found_start = None
        max_start = len(actual_keys) - group.original_length + 1
        for candidate_start in range(search_start, max_start + 1):
            candidate_end = candidate_start + group.original_length - 1
            if actual_keys[candidate_start - 1:candidate_end] == group.original_keys:
                found_start = candidate_start
                break

        if found_start is None:
            raise ValidationError(
                f"{scene_name}: could not locate ruby group starting with {group.original_keys[0]!r}"
            )

        located_groups.append(
            RubyGroup(
                original_start=found_start,
                original_length=group.original_length,
                flat_start=0,
                combined_text=group.combined_text,
                ruby_text=group.ruby_text,
                ruby_count=group.ruby_count,
                original_keys=list(group.original_keys),
                segments=[Segment(segment.type, segment.text, segment.emitted) for segment in group.segments],
            )
        )
        search_start = found_start + group.original_length

    return located_groups


def detect_sheet_layout(worksheet: Any) -> SheetLayout:
    if worksheet.cell(1, 1).value == "Key" and worksheet.cell(1, 2).value == "Value":
        return SheetLayout(key_col=1, translation_col=2, index_col=None, width=2)

    return SheetLayout(key_col=1, translation_col=5, index_col=3, width=5)


def read_sheet_rows(workbook_path: Path) -> tuple[str, SheetLayout, list[list[Any]]]:
    workbook = load_workbook(workbook_path)
    worksheet = workbook[workbook.sheetnames[0]]
    layout = detect_sheet_layout(worksheet)
    rows: list[list[Any]] = []
    for row_idx in range(2, worksheet.max_row + 1):
        rows.append([worksheet.cell(row_idx, column_idx).value for column_idx in range(1, layout.width + 1)])
    sheet_name = worksheet.title
    workbook.close()
    return sheet_name, layout, rows


def validate_keys(expected_rows: list[str], sheet_rows: list[list[Any]], scene_name: str) -> None:
    actual_rows = [row[0] for row in sheet_rows]
    if len(expected_rows) != len(actual_rows):
        raise ValidationError(
            f"{scene_name}: row count mismatch, script has {len(expected_rows)} rows and workbook has {len(actual_rows)} rows"
        )

    for index, (expected, actual) in enumerate(zip(expected_rows, actual_rows), start=2):
        if expected != actual:
            raise ValidationError(
                f"{scene_name}: mismatch at worksheet row {index}: expected {expected!r}, found {actual!r}"
            )


def flatten_sheet_rows(
    sheet_rows: list[list[Any]],
    groups: list[RubyGroup],
    layout: SheetLayout,
) -> list[list[Any]]:
    flattened_rows: list[list[Any]] = []
    group_by_start = {group.original_start: group for group in groups}
    original_idx = 1
    flat_idx = 1

    while original_idx <= len(sheet_rows):
        group = group_by_start.get(original_idx)
        if group is None:
            flattened_rows.append(list(sheet_rows[original_idx - 1]))
            original_idx += 1
            flat_idx += 1
            continue

        group.flat_start = flat_idx
        combined_row = list(sheet_rows[original_idx - 1])
        ruby_row = list(sheet_rows[original_idx])
        combined_row[0] = group.combined_text
        ruby_row[0] = group.ruby_text
        flattened_rows.append(combined_row)
        flattened_rows.append(ruby_row)

        original_idx += group.original_length
        flat_idx += 2

    reindex_rows(flattened_rows, layout.index_col)
    return flattened_rows


def reindex_rows(rows: list[list[Any]], index_col: int | None) -> None:
    if index_col is None:
        return
    for index, row in enumerate(rows):
        row[index_col - 1] = index


def write_sheet_rows(workbook_path: Path, rows: list[list[Any]], layout: SheetLayout) -> None:
    workbook = load_workbook(workbook_path)
    worksheet = workbook[workbook.sheetnames[0]]
    current_count = max(0, worksheet.max_row - 1)
    target_count = len(rows)

    if current_count > target_count:
        worksheet.delete_rows(target_count + 2, current_count - target_count)
    elif current_count < target_count:
        worksheet.insert_rows(current_count + 2, target_count - current_count)

    for row_idx, row in enumerate(rows, start=2):
        for column_idx in range(1, layout.width + 1):
            value = row[column_idx - 1] if column_idx - 1 < len(row) else None
            worksheet.cell(row_idx, column_idx).value = value

    workbook.save(workbook_path)
    workbook.close()


def scene_metadata_path(game_dir: Path, workbook_name: str) -> Path:
    return game_dir / TEXT_DIRNAME / MAP_DIRNAME / f"{workbook_name}.json"


def build_metadata(
    workbook_name: str,
    sheet_name: str,
    original_row_count: int,
    flattened_row_count: int,
    groups: list[RubyGroup],
) -> SceneMetadata:
    return SceneMetadata(
        version=1,
        workbook=workbook_name,
        sheet_name=sheet_name,
        original_row_count=original_row_count,
        flattened_row_count=flattened_row_count,
        ruby_joiner=RUBY_JOINER,
        groups=groups,
    )


def save_metadata(game_dir: Path, metadata: SceneMetadata) -> None:
    map_dir = game_dir / TEXT_DIRNAME / MAP_DIRNAME
    map_dir.mkdir(parents=True, exist_ok=True)
    map_path = scene_metadata_path(game_dir, metadata.workbook)
    payload = asdict(metadata)
    map_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_metadata(game_dir: Path, workbook_name: str) -> SceneMetadata:
    map_path = scene_metadata_path(game_dir, workbook_name)
    if not map_path.exists():
        raise FileNotFoundError(f"Missing metadata for {workbook_name}: {map_path}")

    payload = json.loads(map_path.read_text(encoding="utf-8"))
    groups = [
        RubyGroup(
            original_start=group["original_start"],
            original_length=group["original_length"],
            flat_start=group["flat_start"],
            combined_text=group["combined_text"],
            ruby_text=group["ruby_text"],
            ruby_count=group["ruby_count"],
            original_keys=group["original_keys"],
            segments=[Segment(**segment) for segment in group["segments"]],
        )
        for group in payload["groups"]
    ]
    return SceneMetadata(
        version=payload["version"],
        workbook=payload["workbook"],
        sheet_name=payload["sheet_name"],
        original_row_count=payload["original_row_count"],
        flattened_row_count=payload["flattened_row_count"],
        ruby_joiner=payload["ruby_joiner"],
        groups=groups,
    )


def split_helper_translation(helper_text: str, ruby_count: int) -> list[str]:
    if ruby_count == 0:
        return []
    if not helper_text:
        return [""] * ruby_count

    pieces = [piece.strip() for piece in helper_text.split(RUBY_JOINER)]
    if len(pieces) < ruby_count:
        pieces.extend([""] * (ruby_count - len(pieces)))
    return pieces[:ruby_count]


def split_combined_translation(combined_text: str, base_texts: list[str]) -> list[str] | None:
    plain_chunks: list[str] = []
    cursor = 0

    for base_text in base_texts:
        if not base_text:
            return None
        found_span = find_base_text_span(combined_text, base_text, cursor)
        if found_span is None:
            return None
        found_at, found_end = found_span
        plain_chunks.append(combined_text[cursor:found_at])
        cursor = found_end

    plain_chunks.append(combined_text[cursor:])
    return plain_chunks


def find_base_text_span(combined_text: str, base_text: str, cursor: int) -> tuple[int, int] | None:
    found_at = combined_text.find(base_text, cursor)
    if found_at >= 0:
        return found_at, found_at + len(base_text)

    base_tokens = sorted(match.group(0).lower() for match in TOKEN_RE.finditer(base_text))
    if not base_tokens:
        return None

    token_matches = list(TOKEN_RE.finditer(combined_text[cursor:]))
    best_span: tuple[int, int] | None = None

    for start_index in range(len(token_matches)):
        collected_tokens: list[str] = []
        for end_index in range(start_index, len(token_matches)):
            collected_tokens.append(token_matches[end_index].group(0).lower())
            if len(collected_tokens) > len(base_tokens):
                break
            if sorted(collected_tokens) != base_tokens:
                continue

            start = cursor + token_matches[start_index].start()
            end = cursor + token_matches[end_index].end()
            candidate = (start, end)
            if best_span is None or (candidate[1] - candidate[0]) < (best_span[1] - best_span[0]):
                best_span = candidate
            break

    return best_span


def build_group_translations(
    group: RubyGroup,
    combined_translation: str,
    helper_translation: str,
    workbook_name: str,
) -> list[str]:
    ruby_translations = split_helper_translation(helper_translation, group.ruby_count)
    base_translations = RESTORE_BASE_TEXT_OVERRIDES.get(
        (workbook_name, group.flat_start),
        ruby_translations,
    )
    plain_chunks = split_combined_translation(combined_translation, base_translations)
    if plain_chunks is None:
        plain_chunks = [""] * (group.ruby_count + 1)
        if plain_chunks:
            plain_chunks[0] = combined_translation

    emitted_translations: list[str] = []
    pending_plain = ""
    plain_index = 0
    ruby_index = 0
    base_index = 0

    for segment in group.segments:
        if segment.type == "plain":
            segment_translation = plain_chunks[plain_index]
            plain_index += 1
            if segment.emitted:
                emitted_translations.append(pending_plain + segment_translation)
                pending_plain = ""
            else:
                pending_plain += segment_translation
            continue

        if segment.type == "ruby":
            ruby_translation = ruby_translations[ruby_index] if ruby_index < len(ruby_translations) else ""
            ruby_index += 1
            emitted_translations.append(pending_plain + ruby_translation)
            pending_plain = ""
            continue

        base_translation = base_translations[base_index] if base_index < len(base_translations) else ""
        base_index += 1
        emitted_translations.append(pending_plain + base_translation)
        pending_plain = ""

    if pending_plain:
        if emitted_translations:
            emitted_translations[-1] += pending_plain
        else:
            emitted_translations.append(pending_plain)

    if len(emitted_translations) != group.original_length:
        raise ValidationError(
            f"{group.original_keys[0]!r}: restore produced {len(emitted_translations)} rows, expected {group.original_length}"
        )

    return emitted_translations


def restore_sheet_rows(
    flat_rows: list[list[Any]],
    metadata: SceneMetadata,
    layout: SheetLayout,
) -> list[list[Any]]:
    restored_rows: list[list[Any]] = []
    group_by_flat_start = {group.flat_start: group for group in metadata.groups}
    flat_index = 1

    while flat_index <= len(flat_rows):
        group = group_by_flat_start.get(flat_index)
        if group is None:
            restored_rows.append(list(flat_rows[flat_index - 1]))
            flat_index += 1
            continue

        combined_row = flat_rows[flat_index - 1]
        helper_row = flat_rows[flat_index]
        combined_translation = combined_row[layout.translation_col - 1] or ""
        helper_translation = helper_row[layout.translation_col - 1] or ""
        row_translations = build_group_translations(
            group,
            combined_translation,
            helper_translation,
            metadata.workbook,
        )

        for key_text, translation_text in zip(group.original_keys, row_translations):
            restored_row = [None] * layout.width
            restored_row[layout.key_col - 1] = key_text
            restored_row[layout.translation_col - 1] = translation_text or None
            restored_rows.append(restored_row)

        flat_index += 2

    reindex_rows(restored_rows, layout.index_col)
    return restored_rows


def iter_scene_pairs(game_dir: Path) -> list[tuple[Path, Path]]:
    text_dir = game_dir / TEXT_DIRNAME
    script_dir = game_dir / SCRIPT_DIRNAME
    workbooks = sorted(text_dir.glob("*.xlsx"))
    pairs: list[tuple[Path, Path]] = []

    for workbook_path in workbooks:
        scene_path = script_dir / workbook_path.name[:-5]
        if not scene_path.exists():
            raise FileNotFoundError(f"Missing script for workbook {workbook_path.name}: {scene_path}")
        pairs.append((scene_path, workbook_path))

    return pairs


def flatten_game_dir(game_dir: Path) -> None:
    for scene_path, workbook_path in iter_scene_pairs(game_dir):
        _, groups = parse_scene(scene_path)
        if not groups:
            continue
        sheet_name, layout, sheet_rows = read_sheet_rows(workbook_path)
        located_groups = locate_ruby_groups(sheet_rows, groups, workbook_path.name)
        flattened_rows = flatten_sheet_rows(sheet_rows, located_groups, layout)
        metadata = build_metadata(
            workbook_name=workbook_path.name,
            sheet_name=sheet_name,
            original_row_count=len(sheet_rows),
            flattened_row_count=len(flattened_rows),
            groups=located_groups,
        )
        write_sheet_rows(workbook_path, flattened_rows, layout)
        save_metadata(game_dir, metadata)
        print(f"flattened {workbook_path.name}: {len(sheet_rows)} -> {len(flattened_rows)} rows")


def restore_game_dir(game_dir: Path, workbook_dirname: str, metadata_dirname: str) -> None:
    map_dir = game_dir / metadata_dirname / MAP_DIRNAME
    for map_path in sorted(map_dir.glob("*.json")):
        workbook_path = game_dir / workbook_dirname / map_path.stem
        metadata = load_metadata(game_dir, workbook_path.name)
        _, layout, flat_rows = read_sheet_rows(workbook_path)
        if len(flat_rows) != metadata.flattened_row_count:
            raise ValidationError(
                f"{workbook_path.name}: flattened row count changed, expected {metadata.flattened_row_count} rows and found {len(flat_rows)}"
            )
        restored_rows = restore_sheet_rows(flat_rows, metadata, layout)
        write_sheet_rows(workbook_path, restored_rows, layout)
        print(f"restored {workbook_path.name}: {len(flat_rows)} -> {len(restored_rows)} rows")


def validate_game_dir(game_dir: Path) -> None:
    for scene_path, workbook_path in iter_scene_pairs(game_dir):
        _, groups = parse_scene(scene_path)
        if not groups:
            continue
        _, _, sheet_rows = read_sheet_rows(workbook_path)
        locate_ruby_groups(sheet_rows, groups, workbook_path.name)
        print(f"validated {workbook_path.name}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Flatten and restore ruby-heavy XLSX files for Chicchakunai Mon.",
    )
    parser.add_argument(
        "command",
        choices=["flatten", "restore", "validate"],
        help="Operation to run on the Text folder.",
    )
    parser.add_argument(
        "--game-dir",
        type=Path,
        default=DEFAULT_GAME_DIR,
        help="Path to the Chicchakunai Mon game folder.",
    )
    parser.add_argument(
        "--workbook-dirname",
        default=TEXT_DIRNAME,
        help="Workbook subdirectory to operate on for restore. Defaults to Text.",
    )
    parser.add_argument(
        "--metadata-dirname",
        default=TEXT_DIRNAME,
        help="Subdirectory that contains .ruby_flatten_maps. Defaults to Text.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    game_dir = args.game_dir.resolve()

    try:
        if args.command == "flatten":
            flatten_game_dir(game_dir)
        elif args.command == "restore":
            restore_game_dir(game_dir, args.workbook_dirname, args.metadata_dirname)
        else:
            validate_game_dir(game_dir)
    except (FileNotFoundError, ValidationError) as exc:
        print(str(exc), file=sys.stderr)
        return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(main())